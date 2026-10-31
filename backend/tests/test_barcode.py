import io

import pandas as pd
import pytest

from app.barcode import normalize_barcode, normalize_excel_barcode

from .helpers import make_location, make_product, make_warehouse


# ---------------------------------------------------------------- التنظيف والتحقق
@pytest.mark.parametrize("raw,expected", [
    (None, None), ("", None), ("   ", None), ("000123", "000123"), ("  AbC-12  ", "AbC-12"), ("6221234567890", "6221234567890"),
    ("a", "a"), ("A" * 64, "A" * 64), ("x/y_z.1", "x/y_z.1"),
])
def test_normalize_valid(raw, expected):
    assert normalize_barcode(raw) == expected


@pytest.mark.parametrize("raw", ["A" * 65, "ab cd", "باركود", "=SUM(A1)", "-123", "+123", "@x", "\tabc d", "ab\u200bcd"])
def test_normalize_rejected(raw):
    with pytest.raises(ValueError):
        normalize_barcode(raw)


@pytest.mark.parametrize("raw,expected", [("1234567890123.0", "1234567890123"), ("000123", "000123"), ("nan", None), (None, None), ("12.50", "12.50")])
def test_excel_normalize(raw, expected):
    assert normalize_excel_barcode(raw) == expected


def test_excel_scientific_rejected():
    with pytest.raises(ValueError, match="كنص"):
        normalize_excel_barcode("1.23E+12")


# ---------------------------------------------------------------- API
def _create(c, name="م", barcode=None, sku=None, **kw):
    body = {"name": name, "price": 1, **kw}
    if barcode is not None:
        body["barcode"] = barcode
    if sku is not None:
        body["sku"] = sku
    return c.post("/api/products", json=body)


def test_create_and_update_with_barcode(admin):
    r = _create(admin, "منتج", barcode="  000123 ")
    assert r.status_code == 200 and r.json()["barcode"] == "000123"
    pid = r.json()["id"]
    assert admin.put(f"/api/products/{pid}", json={"name": "جديد"}).json()["barcode"] == "000123"  # غير مرسل = من غير تغيير
    assert admin.put(f"/api/products/{pid}", json={"barcode": "ABC9"}).json()["barcode"] == "ABC9"
    assert admin.put(f"/api/products/{pid}", json={"barcode": ""}).json()["barcode"] is None  # مسح
    assert _create(admin, "من غير باركود").json()["barcode"] is None
    assert _create(admin, "null", barcode=None).status_code == 200


@pytest.mark.parametrize("bad", ["a b", "=1+1", "-5", "باركود", "A" * 65])
def test_create_rejects_invalid_barcode(admin, bad):
    assert _create(admin, "x", barcode=bad).status_code == 422


def test_uniqueness_is_case_insensitive(admin):
    assert _create(admin, "أ", barcode="Abc123").status_code == 200
    r = _create(admin, "ب", barcode="aBC123")
    assert r.status_code == 400 and "مستخدم" in r.json()["detail"]
    other = _create(admin, "ج", barcode="Zzz1").json()["id"]
    assert admin.put(f"/api/products/{other}", json={"barcode": "ABC123"}).status_code == 400
    assert admin.put(f"/api/products/{other}", json={"barcode": "ZZZ1"}).status_code == 200  # نفس منتجه: مسموح


def test_barcode_cannot_equal_other_products_sku_and_vice_versa(admin):
    assert _create(admin, "أ", sku="SKU-1").status_code == 200
    r = _create(admin, "ب", barcode="sku-1")
    assert r.status_code == 400 and "SKU" in r.json()["detail"]
    assert _create(admin, "ج", barcode="BAR-9").status_code == 200
    r = _create(admin, "د", sku="bar-9")
    assert r.status_code == 400 and "باركود" in r.json()["detail"]
    # مسموح: SKU = الباركود لنفس المنتج
    assert _create(admin, "هـ", sku="SAME-1", barcode="SAME-1").status_code == 200


def test_search_by_barcode_sku_and_name(admin):
    a = _create(admin, "كابل شحن", barcode="6221111111111", sku="CBL-1").json()
    b = _create(admin, "ماوس", barcode="6222222222222").json()
    assert [p["id"] for p in admin.get("/api/products", params={"search": "6221111", "page": 1}).json()["items"]] == [a["id"]]
    lk = admin.get("/api/products/lookup", params={"q": "6222222222222"}).json()
    assert [p["id"] for p in lk] == [b["id"]] and lk[0]["barcode"] == "6222222222222"
    assert admin.get("/api/products/lookup", params={"q": "CBL-1"}).json()[0]["id"] == a["id"]
    assert admin.get("/api/products/lookup", params={"q": "ماوس"}).json()[0]["id"] == b["id"]


def test_lookup_exact_barcode_match_comes_first(admin):
    _create(admin, "منتج 1234 مشابه", barcode="X1234567")
    exact = _create(admin, "zzz آخر", barcode="1234").json()
    assert admin.get("/api/products/lookup", params={"q": "1234"}).json()[0]["id"] == exact["id"]


def test_by_barcode_endpoint(admin, login):
    wh = make_warehouse(admin)
    loc = make_location(admin, warehouse_id=wh["id"], code="A1")
    p = make_product(admin, "صنف", barcode="BC-100", locations=[(loc["id"], 7)])
    r = admin.get("/api/products/by-barcode/bc-100")  # case-insensitive
    assert r.status_code == 200
    body = r.json()
    assert body["id"] == p["id"] and body["barcode"] == "BC-100"
    assert body["locations"] == [{"location_id": loc["id"], "full_code": "W1-A1", "quantity": 7}]
    sku_p = _create(admin, "بـ SKU", sku="S-77").json()
    assert admin.get("/api/products/by-barcode/S-77").json()["id"] == sku_p["id"]  # SKU كبديل
    assert admin.get("/api/products/by-barcode/nope").status_code == 404
    admin.patch(f"/api/products/{p['id']}/active", json={"is_active": False})
    r = admin.get("/api/products/by-barcode/BC-100")
    assert r.status_code == 409 and "معطّل" in r.json()["detail"]
    assert login("noperm", []).get("/api/products/by-barcode/BC-100").status_code == 403


def test_by_barcode_prefers_barcode_over_sku(admin):
    # باركود منتج أ و SKU منتج ب مينفعش يتساووا أصلًا (قفل التضارب)، فالتطابق دايمًا واحد
    a = _create(admin, "أ", barcode="DUP-1").json()
    assert _create(admin, "ب", sku="DUP-1").status_code == 400
    assert admin.get("/api/products/by-barcode/DUP-1").json()["id"] == a["id"]


# ---------------------------------------------------------------- Excel
def _xlsx(rows, columns=None):
    buf = io.BytesIO()
    pd.DataFrame(rows, columns=columns).to_excel(buf, index=False)
    return buf.getvalue()


def _import(c, rows, columns=None):
    return c.post("/api/excel/products/import", files={"file": ("p.xlsx", _xlsx(rows, columns), "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")}).json()


def _products(c):
    return {p["name"]: p for p in c.get("/api/products", params={"page": 1, "page_size": 100}).json()["items"]}


def test_import_old_file_without_barcode_column_keeps_existing_barcode(admin):
    p = _create(admin, "قديم", barcode="KEEP-1").json()
    out = _import(admin, [{"ID": p["id"], "الاسم": "قديم", "السعر": 5}])
    assert out["updated"] == 1 and out["errors"] == []
    assert _products(admin)["قديم"]["barcode"] == "KEEP-1"


def test_import_with_barcode_creates_updates_and_empty_cell_does_not_clear(admin):
    out = _import(admin, [{"الاسم": "جديد", "الباركود": "000987"}])
    assert out["created"] == 1 and _products(admin)["جديد"]["barcode"] == "000987"
    pid = _products(admin)["جديد"]["id"]
    out = _import(admin, [{"ID": pid, "الاسم": "جديد", "الباركود": None, "السعر": 9}])
    assert out["updated"] == 1 and _products(admin)["جديد"]["barcode"] == "000987"  # الخلية الفاضية ما بتمسحش
    out = _import(admin, [{"الاسم": "اسم تاني", "الباركود": "000987", "السعر": 3}])  # مطابقة بالباركود لمنتج موجود
    assert out["updated"] == 1 and out["created"] == 0
    assert _products(admin)["اسم تاني"]["price"] == 3
    assert _products(admin)["اسم تاني"]["barcode"] == "000987" and "جديد" not in _products(admin)


def test_import_float_and_scientific_barcodes(admin):
    buf = io.BytesIO()
    pd.DataFrame([{"الاسم": "رقمي", "الباركود": 6221234567890}, {"الاسم": "علمي", "الباركود": "1.23E+12"}]).to_excel(buf, index=False)
    out = admin.post("/api/excel/products/import", files={"file": ("p.xlsx", buf.getvalue(), "application/octet-stream")}).json()
    assert out["created"] == 1 and any("كنص" in e for e in out["errors"])
    assert _products(admin)["رقمي"]["barcode"] == "6221234567890"


def test_import_conflicts_and_duplicates(admin):
    _create(admin, "موجود", barcode="TAKEN-1", sku="S-1")
    _create(admin, "تاني", sku="S-2")
    # SKU بيطابق منتج "تاني" والباركود تابع لمنتج "موجود" => رفض الصف
    out = _import(admin, [{"الاسم": "تاني", "SKU": "S-2", "الباركود": "TAKEN-1"}])
    assert out["updated"] == 0 and any("تابع لمنتج" in e for e in out["errors"])
    assert _products(admin)["تاني"]["barcode"] is None
    out = _import(admin, [{"الاسم": "ب", "الباركود": "DUP-9"}, {"الاسم": "ج", "الباركود": "dup-9"}])
    assert out["created"] == 1 and any("منتجين مختلفين" in e for e in out["errors"])
    out = _import(admin, [{"الاسم": "د", "SKU": "TAKEN-1"}])
    assert out["created"] == 0 and any("كباركود" in e for e in out["errors"])


def test_export_has_text_barcode_column_and_roundtrip_keeps_zeros(admin):
    _create(admin, "بأصفار", barcode="000123")
    r = admin.get("/api/excel/products/export")
    df = pd.read_excel(io.BytesIO(r.content), dtype=str)
    assert list(df.columns)[:3] == ["ID", "SKU", "الباركود"] and df.loc[0, "الباركود"] == "000123"
    from openpyxl import load_workbook

    ws = load_workbook(io.BytesIO(r.content))["المنتجات"]
    assert ws["C2"].number_format == "@" and ws["C2"].value == "000123"
    out = admin.post("/api/excel/products/import", files={"file": ("p.xlsx", r.content, "application/octet-stream")}).json()
    assert out["errors"] == [] and _products(admin)["بأصفار"]["barcode"] == "000123"


def test_template_has_barcode_column_and_instructions(admin):
    r = admin.get("/api/excel/products/template")
    sheets = pd.read_excel(io.BytesIO(r.content), sheet_name=None, dtype=str)
    assert "الباركود" in sheets["المنتجات"].columns
    assert any("الباركود" in str(x) for x in sheets["تعليمات"]["تعليمات"])
